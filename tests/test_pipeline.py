"""Tests for the rules that decide what reaches the sales team.

The scrapers are not tested here - they run against live sites and are
calibrated by `probe`. What is tested is everything that turns raw creditor
rows into a prospect list, because that is where a silent bug publishes a
wrong name, a wrong amount, or an existing client as a new lead.
"""

from __future__ import annotations

import sys
from datetime import date
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "src"))

import pytest

from creditor_sourcing import aggregate, ledger, qualify
from creditor_sourcing.models import Creditor, Matter, normalise_name
from creditor_sourcing.parse import creditor_tables
from creditor_sourcing.parse.creditor_tables import parse_lines
from creditor_sourcing.sources import (
    abn_lookup,
    asic_connect,
    asic_dataset,
    worrells,
)
from creditor_sourcing.sources.worrells import details_view_url


def creditor(name, debtor="Bust Co Pty Ltd", matter="m1", amount=50000.0, **kw):
    return Creditor(name, debtor, matter, amount, **kw)


class TestNormalisation:
    @pytest.mark.parametrize(
        "a,b",
        [
            ("Acme Building Supplies Pty Ltd", "ACME BUILDING SUPPLIES"),
            ("Riverside Timber & Hardware", "Riverside Timber and Hardware Pty Ltd"),
            # Corporate-family words fold too - "Acme Group" and "Acme Australia"
            # are the same customer relationship for prospecting purposes.
            ("The Steel Group Australia Limited", "Steel  Pty. Ltd."),
        ],
    )
    def test_variants_fold_together(self, a, b):
        assert normalise_name(a) == normalise_name(b)

    def test_distinct_companies_stay_distinct(self):
        assert normalise_name("Acme Timber") != normalise_name("Acme Steel")


class TestAggregation:
    def test_same_company_across_matters_is_one_prospect(self):
        prospects = aggregate.build(
            [
                creditor("Acme Pty Ltd", "Bust Co", "m1", 42000.0),
                creditor("ACME", "Other Bust", "m2", 18000.0),
            ]
        )
        assert len(prospects) == 1
        assert prospects[0].matter_count == 2
        assert prospects[0].total_exposure_aud == 60000.0

    def test_related_parties_never_become_prospects(self):
        assert aggregate.build([creditor("Smith Family Trust", related_party=True)]) == []


class TestQualification:
    def test_under_the_floor_is_dropped(self):
        [p] = qualify.apply(aggregate.build([creditor("Tiny Widgets Pty Ltd", amount=4999.0)]))
        assert not p.qualified and "under the" in p.disqualified_reason

    def test_at_the_floor_is_kept(self):
        [p] = qualify.apply(aggregate.build([creditor("Small Widgets Pty Ltd", amount=5000.0)]))
        assert p.qualified

    @pytest.mark.parametrize(
        "name",
        [
            "Australian Taxation Office",
            "Westpac Banking Corporation",
            "Worrells Solvency & Forensic Accountants",
            "QBE Insurance (Australia) Limited",
            "Smith & Co Lawyers",
            "Employee Entitlements",
        ],
    )
    def test_non_trade_creditors_are_dropped(self, name):
        [p] = qualify.apply(aggregate.build([creditor(name, amount=250000.0)]))
        assert not p.qualified, f"{name} should not be a prospect"

    def test_ordinary_supplier_survives(self):
        [p] = qualify.apply(aggregate.build([creditor("Riverside Timber Supplies Pty Ltd")]))
        assert p.qualified

    def test_existing_client_is_dropped_despite_spelling_drift(self):
        [p] = qualify.apply(
            aggregate.build([creditor("Riverside Timber & Hardware Pty Ltd", amount=250000.0)]),
            {"riverside timber hardware": "Riverside Timber and Hardware Pty Ltd"},
        )
        assert not p.qualified and "Existing NCI client" in p.disqualified_reason

    def test_word_boundaries_stop_false_positives(self):
        # "agl" is an excluded utility; "Eagle" must not match it.
        [p] = qualify.apply(aggregate.build([creditor("Eagle Fasteners Pty Ltd")]))
        assert p.qualified


class TestScoring:
    def test_repeat_exposure_outranks_a_bigger_single_debt(self):
        repeat = aggregate.build(
            [creditor("Repeat Co", "A", "m1", 30000.0), creditor("Repeat Co", "B", "m2", 30000.0)]
        )[0]
        once = aggregate.build([creditor("Single Co", "A", "m1", 400000.0)])[0]
        assert qualify.score(repeat) > qualify.score(once)

    def test_score_is_bounded(self):
        huge = aggregate.build(
            [creditor("Big Co", f"D{i}", f"m{i}", 5_000_000.0) for i in range(10)]
        )[0]
        assert 0 <= qualify.score(huge) <= 100


class TestCreditorTableParsing:
    LINES = [
        "Listing of known creditors",
        "Name Address Related Party Amount",
        "Acme Building Supplies Pty Ltd 12 Industry Rd Dandenong VIC No 42,500.00",
        "Smith Family Trust 3 Hill St Toorak VIC Yes 60,000.00",
        "Contents ................................ 4",
        "Total 230,500.00",
        "Page 7",
    ]

    def test_reads_real_rows_only(self):
        rows = parse_lines(self.LINES, "Bust Co", "m1", "asic")
        assert [r.creditor_name for r in rows] == [
            "Acme Building Supplies Pty Ltd",
            "Smith Family Trust",
        ]

    def test_amount_and_address_split(self):
        row = parse_lines(self.LINES, "Bust Co", "m1", "asic")[0]
        assert row.amount_aud == 42500.0
        assert row.address == "12 Industry Rd Dandenong VIC"

    def test_related_party_flag(self):
        assert parse_lines(self.LINES, "Bust Co", "m1", "asic")[1].related_party

    def test_contents_lines_are_never_creditors(self):
        rows = parse_lines(["Schedule of debts .......... 12"], "Bust Co", "m1", "asic")
        assert rows == []


class TestWorrellsDetails:
    """The File Details panel on a DetailsView page.

    DETAIL_HTML is the verbatim visible text of the live page for
    MAYDE ELECTRICAL PTY LTD, captured by the diagnose workflow. The panel is
    label/value pairs in inline elements, not a table.
    """

    DETAIL_HTML = (
        "<html><body>File Information MAYDE ELECTRICAL PTY LTD File Details "
        "ACN/Estate#: 168 364 298 Office Name: North Lakes Trading Name None known "
        "Principal Lee Crosthwaite Admin Type Creditors Vol Manager: Broderick Dipple "
        "Start Date 11/09/2026 Exec Analyst Broderick Dipple Status Priority "
        "Contact Person Joseph Eckersley Industry Building/Construction "
        "Appointee Andrew Worrell I want to Lodge a Proof of Debt</body></html>"
    )

    def test_values_stop_at_the_next_panel_label(self):
        # Without every panel label as a boundary - including the ones we do
        # not map - Start Date swallowed the rest of the panel.
        details = worrells.parse_details(self.DETAIL_HTML)
        assert details["appointment_date"] == "11/09/2026"
        assert details["appointment_type"] == "Creditors Vol"
        assert details["office"] == "North Lakes"

    def test_acn_is_extracted_and_padded(self):
        # The ACN is what lets a Worrells matter reconcile against the same
        # company in the ASIC workbook.
        matter = worrells.apply_details(
            Matter(source="worrells", company_name="MAYDE ELECTRICAL PTY LTD"),
            self.DETAIL_HTML,
        )
        assert matter.acn == "168364298"

    def test_start_date_becomes_an_iso_date(self):
        matter = worrells.apply_details(
            Matter(source="worrells", company_name="X"), self.DETAIL_HTML)
        assert matter.appointment_date == "2026-09-11"

    def test_industry_and_appointee_are_captured(self):
        matter = worrells.apply_details(
            Matter(source="worrells", company_name="X"), self.DETAIL_HTML)
        assert matter.industry == "Building/Construction"
        assert matter.practitioner == "Andrew Worrell"
        assert matter.practitioner_firm == "Worrells"

    def test_placeholder_values_are_dropped(self):
        assert "trading_name" not in worrells.parse_details(self.DETAIL_HTML)

    def test_existing_values_are_never_overwritten(self):
        matter = Matter(source="worrells", company_name="X", acn="999888777",
                        industry="Manufacturing")
        worrells.apply_details(matter, self.DETAIL_HTML)
        assert matter.acn == "999888777"
        assert matter.industry == "Manufacturing"

    def test_a_bare_page_yields_nothing_rather_than_junk(self):
        assert worrells.parse_details("<html><body>Nothing here</body></html>") == {}


class TestWorrellsDocuments:
    """Document links on a DetailsView page.

    Confirmed live: most matters carry no documents (they are days old), while
    a matured one such as Valera Recycling Pty Ltd carries "First Advice" and
    "2nd Advice" - both creditor listings, downloadable unauthenticated.
    """

    HTML = """<html><body>
      <a href="/">home</a>
      <a href="/FileInformation">File Information</a>
      <a href="/WebDocuments/12345/first-advice.pdf">First Advice</a>
      <a href="/WebDocuments/12345/2nd-advice.pdf">2nd Advice</a>
      <a href="/WebDocuments/12345/remuneration.pdf">Remuneration Report</a>
      <a href="/Privacy">Privacy</a>
    </body></html>"""

    BASE = "https://customerportal.worrells.net.au"

    def test_only_pdfs_are_returned(self):
        docs = worrells.parse_documents(self.HTML, self.BASE)
        assert len(docs) == 3
        assert all(d["url"].endswith(".pdf") for d in docs)

    def test_creditor_documents_rank_first(self):
        names = [d["name"] for d in worrells.parse_documents(self.HTML, self.BASE)]
        assert names[:2] == ["2nd Advice", "First Advice"]
        assert "Remuneration Report" not in names[:2]

    def test_ranking_follows_measured_yield(self):
        # Across 14 published documents: Initial Advice carried a listing 4/4,
        # 2nd Advice 2/3, First Advice only 1/6. So the order is not
        # chronological - it is by how often the document actually has the
        # creditor annexure.
        html = """<html><body>
          <a href="/WebDocuments/1/first.pdf">First Advice</a>
          <a href="/WebDocuments/1/second.pdf">2nd Advice</a>
          <a href="/WebDocuments/1/initial.pdf">Initial Advice</a>
          <a href="/WebDocuments/1/plan.pdf">Notice of Plan</a>
        </body></html>"""
        names = [d["name"] for d in worrells.creditor_documents(html, self.BASE)]
        assert names == ["Notice of Plan", "Initial Advice", "2nd Advice",
                         "First Advice"]

    def test_creditor_documents_excludes_other_reports(self):
        names = [d["name"] for d in worrells.creditor_documents(self.HTML, self.BASE)]
        assert names == ["2nd Advice", "First Advice"]

    def test_urls_are_absolute(self):
        doc = worrells.creditor_documents(self.HTML, self.BASE)[0]
        assert doc["url"] == f"{self.BASE}/WebDocuments/12345/2nd-advice.pdf"

    def test_a_new_matter_with_no_documents_is_not_an_error(self):
        bare = """<html><body><a href="/">h</a><a href="/Privacy">p</a></body></html>"""
        assert worrells.parse_documents(bare, self.BASE) == []


class TestWorrellsUrl:
    LIST_URL = (
        "https://customerportal.worrells.net.au/FileInformation/FileInformation/"
        "202502200334422382761DFEC99D43BE"
    )
    VIEW_URL = (
        "https://customerportal.worrells.net.au/FileInformation/"
        "FileInformationDetailsView/202502200334422382761DFEC99D43BE"
    )

    def test_list_url_is_rewritten_to_the_documents_page(self):
        assert details_view_url(self.LIST_URL) == self.VIEW_URL

    def test_already_correct_url_is_unchanged(self):
        assert details_view_url(self.VIEW_URL) == self.VIEW_URL

    def test_bare_id_is_accepted(self):
        assert details_view_url("202502200334422382761DFEC99D43BE") == self.VIEW_URL

    def test_garbage_is_rejected(self):
        with pytest.raises(ValueError):
            details_view_url("https://customerportal.worrells.net.au/FileInformation")


class TestAsicDataSet:
    """The "Data set" sheet of ASIC's insolvency statistics workbook.

    HEADER and the sheet shape below are ASIC's real ones, captured from the
    published workbook by the "ASIC data set schema" workflow: five blank/title
    rows, a merged group-label row, then the header on row 7.
    """

    # Exactly as published, trailing spaces and embedded newlines included.
    HEADER = [
        "Data to", "ACN No ", "Organisation name ", "Appointee (person or company)",
        "Appointment type", "Effective date", "Period\n(Year month)",
        "Period (financial year)", "Industry type (division)",
        "Industry type (subdivision)", "Industry type (group)",
        "State of incorporation (state or territory)",
        "Principal place of business (state or territory)",
        "Principal place of business (area)",
        "Principal place of business (postcode)",
        "Series 1 \n(companies entering)", "Series 2 \n(all appointments)",
    ]

    @staticmethod
    def row(name, acn, appointment, when, series_1=1, industry="Construction",
            subdivision="Building Construction", state="Victoria", postcode=3000):
        return [None, acn, name, "JONES, MICHAEL GREGORY", appointment, when,
                "2026 09", "2026-2027", industry, subdivision, "Group",
                "Victoria", state, "Melbourne - West", postcode, series_1, 1]

    ROWS = [
        row.__func__("TANCRED BROTHERS PTY LTD", 25712, "Court liquidation",
                     date(2026, 9, 3), industry="Retail Trade",
                     subdivision="Food Retailing", state="Queensland", postcode=4305),
        row.__func__("JINDONG HOLDINGS PTY LTD", 612927974,
                     "Creditors' voluntary liquidation", date(2026, 9, 8)),
    ]

    @classmethod
    def workbook(cls, header=None, rows=None, sheet_name="Data set"):
        import io

        from openpyxl import Workbook

        wb = Workbook()
        wb.remove(wb.active)
        wb.create_sheet("Contents")
        wb.create_sheet("1.1")
        ws = wb.create_sheet(sheet_name)
        for _ in range(5):
            ws.append([])
        ws.append([None, "Organisation details", None, "Appointee", "Role"])
        ws.append(header or cls.HEADER)
        for row in rows if rows is not None else cls.ROWS:
            ws.append(row)
        buf = io.BytesIO()
        wb.save(buf)
        return buf.getvalue()

    def parse(self, **kw):
        return asic_dataset.parse(self.workbook(**kw), lookback_days=30)

    def test_header_is_found_on_row_seven(self):
        assert [m.company_name for m in self.parse()] == [
            "TANCRED BROTHERS PTY LTD", "JINDONG HOLDINGS PTY LTD",
        ]

    def test_low_acn_is_zero_padded_not_dropped(self):
        # ACNs are stored as numbers, so 000 025 712 arrives as 25712.
        # Dropping short values loses every early-registered company.
        assert self.parse()[0].acn == "000025712"

    def test_normal_acn_is_unchanged(self):
        assert self.parse()[1].acn == "612927974"

    def test_industry_and_location_are_captured(self):
        matter = self.parse()[0]
        assert (matter.industry, matter.industry_subdivision) == (
            "Retail Trade", "Food Retailing",
        )
        assert (matter.state, matter.postcode) == ("Queensland", "4305")

    def test_practitioner_is_captured(self):
        assert self.parse()[0].practitioner == "JONES, MICHAEL GREGORY"

    def test_repeat_appointments_are_excluded(self):
        # A Series 2 row without the Series 1 flag is a subsequent appointment
        # for a company already counted - it would duplicate the matter.
        rows = [*self.ROWS, self.row("JINDONG HOLDINGS PTY LTD", 612927974,
                                     "Court liquidation", date(2026, 9, 9),
                                     series_1=None)]
        assert len(self.parse(rows=rows)) == 2

    def test_solvent_wind_ups_are_excluded(self):
        # Members' voluntary liquidation is solvent - nobody lost money.
        rows = [*self.ROWS, self.row("RICH OWNERS PTY LTD", 111222333,
                                     "Members' voluntary liquidation",
                                     date(2026, 9, 5))]
        names = [m.company_name for m in self.parse(rows=rows)]
        assert "RICH OWNERS PTY LTD" not in names

    def test_lookback_filters_the_history(self):
        # The real sheet holds 70,000+ appointments back to July 2021.
        rows = [*self.ROWS, self.row("ANCIENT PTY LTD", 999888777,
                                     "Court liquidation", date(2021, 7, 1))]
        assert len(self.parse(rows=rows)) == 2
        assert len(asic_dataset.parse(self.workbook(rows=rows), lookback_days=0)) == 3

    def test_a_renamed_schema_raises_rather_than_returning_nothing(self):
        # "No insolvencies this week" and "the schema moved" must never look
        # the same, or a silent zero gets reported as a quiet week.
        with pytest.raises(RuntimeError, match="company-name column"):
            self.parse(header=["Widget", "Sprocket", "Gizmo"],
                       rows=[["a", "b", "c"]])

    def test_missing_sheet_names_what_was_there(self):
        with pytest.raises(RuntimeError, match="No data-set sheet"):
            self.parse(sheet_name="Summary")

    def test_stale_file_reports_the_lag_instead_of_looking_like_a_quiet_week(self, caplog):
        # Measured on the real 7 Sep 2026 release: the newest appointment was
        # 19 days old, so a 14-day window returned zero and the weekly run
        # would have reported "no insolvencies" rather than "file is stale".
        import logging

        old = date.today().replace(year=date.today().year - 1).isoformat()
        rows = [self.row("STALE CO PTY LTD", 123456789, "Court liquidation", old)]
        with caplog.at_level(logging.ERROR):
            assert asic_dataset.parse(self.workbook(rows=rows), lookback_days=14) == []
        assert "newest appointment in the file" in caplog.text
        assert "lookback window is only 14 days" in caplog.text


class TestCreditorTableGate:
    """Which pages count as a creditor table.

    These are regression tests for a real failure: on live Worrells reports the
    parser returned five "creditors" all named "Fees:" totalling $66,960,
    harvested from the practitioner's remuneration schedule. The reports run
    20-40 pages and mention creditors throughout, so a heading match locates
    narrative far more often than a table.
    """

    # Verbatim shape of a remuneration page from RFCVIC PTY LTD's report:
    # narrative prose with one inline amount, no Related Party column.
    REMUNERATION_PAGE = "\n".join([
        "NOTICE OF PROPOSAL TO CREDITORS",
        "RFCVIC PTY LTD (In Liquidation)",
        "Proposal 2: Future fee approval",
        "That the remuneration of the Liquidator from 10 September 2026 to the",
        "finalisation of the liquidation is determined at a sum equal to the costs",
        "report to creditors of 10 September 2026 in the amount of $31,500.00",
        "Liquidator's remuneration is paid in priority to unsecured creditors",
        "You may vote Yes, No or Object to the proposal being resolved",
        "Fees: 3,405.70",
        "Fees: 11,000.00",
    ])

    # A real creditor listing: the Related Party column gives standalone
    # Yes/No cells, and amounts sit in their own cells.
    LISTING_PAGE = "\n".join([
        "Listing of known creditors",
        "Name", "Address", "Related Party", "Amount",
        "Acme Building Supplies Pty Ltd", "12 Industry Rd Dandenong VIC",
        "No", "42,500.00",
        "Riverside Timber Pty Ltd", "8 Kembla St Wollongong NSW",
        "No", "128,000.00",
    ])

    def test_a_remuneration_page_is_not_a_creditor_table(self):
        assert not creditor_tables.is_creditor_table(self.REMUNERATION_PAGE)

    def test_a_listing_page_is_a_creditor_table(self):
        assert creditor_tables.is_creditor_table(self.LISTING_PAGE)

    def test_narrative_yes_no_does_not_qualify_a_page(self):
        # "You may vote Yes, No or Object" must not count as the Related Party
        # column - only a cell that is nothing but Yes or No does.
        page = "\n".join(["You may vote Yes, No or Object", "1,000.00", "2,000.00"])
        assert not creditor_tables.is_creditor_table(page)

    @pytest.mark.parametrize(
        "line",
        [
            "Fees: 3,405.70",
            "Fees 11,000.00",
            "GST 1,100.00",
            "Total 230,500.00",
            "Remuneration 31,500.00",
            "Disbursements 265.84",
            "Estimated surplus 5,000.00",
        ],
    )
    def test_fee_and_position_labels_are_never_creditors(self, line):
        rows = creditor_tables.parse_lines([line], "Bust Co", "m1", "worrells")
        assert rows == [], f"{line!r} must not become a creditor"

    def test_a_sentence_ending_in_an_amount_is_not_a_creditor(self):
        line = ("report to creditors of 10 September 2026 in the amount of "
                "31,500.00")
        assert creditor_tables.parse_lines([line], "Bust Co", "m1", "worrells") == []

    def test_a_real_supplier_row_still_parses(self):
        line = "Acme Building Supplies Pty Ltd 12 Industry Rd Dandenong VIC No 42,500.00"
        rows = creditor_tables.parse_lines([line], "Bust Co", "m1", "worrells")
        assert len(rows) == 1
        assert rows[0].creditor_name == "Acme Building Supplies Pty Ltd"
        assert rows[0].amount_aud == 42500.0


class TestCellPerLineListing:
    """PyMuPDF emits one cell per line for a ruled table.

    This is the layout a real creditor listing arrives in, and the original
    line-based parser could not read it at all - which is why the only rows it
    ever produced came from narrative and fee lines.
    """

    PAGE = "\n".join([
        "Annexure C",
        "Listing of known creditors",
        "Name", "Address", "Related Party", "Amount",
        "Acme Building Supplies Pty Ltd", "12 Industry Rd Dandenong VIC",
        "No", "42,500.00",
        "Smith Family Trust", "3 Hill St Toorak VIC", "Yes", "60,000.00",
        "Riverside Timber Pty Ltd", "8 Kembla St Wollongong NSW",
        "No", "128,000.00",
    ])

    def rows(self):
        return creditor_tables.parse_cells(
            self.PAGE.splitlines(), "Bust Co Pty Ltd", "m1", "worrells")

    def test_every_row_is_read(self):
        assert [r.creditor_name for r in self.rows()] == [
            "Acme Building Supplies Pty Ltd",
            "Smith Family Trust",
            "Riverside Timber Pty Ltd",
        ]

    def test_the_section_heading_is_not_a_creditor(self):
        assert "Listing of known creditors" not in [
            r.creditor_name for r in self.rows()]

    def test_the_annexure_label_is_not_a_creditor(self):
        assert "Annexure C" not in [r.creditor_name for r in self.rows()]

    def test_column_headers_are_not_creditors(self):
        names = [r.creditor_name for r in self.rows()]
        assert not {"Name", "Address", "Related Party", "Amount"} & set(names)

    def test_amounts_and_addresses_pair_with_the_right_name(self):
        rows = {r.creditor_name: r for r in self.rows()}
        acme = rows["Acme Building Supplies Pty Ltd"]
        assert acme.amount_aud == 42500.0
        assert acme.address == "12 Industry Rd Dandenong VIC"
        assert rows["Riverside Timber Pty Ltd"].amount_aud == 128000.0

    def test_related_party_flag_follows_the_yes_no_cell(self):
        rows = {r.creditor_name: r for r in self.rows()}
        assert rows["Smith Family Trust"].related_party is True
        assert rows["Acme Building Supplies Pty Ltd"].related_party is False

    def test_an_extra_creditor_type_column_does_not_shift_the_data(self):
        # Some practitioners add a Creditor Type column; anchoring on the
        # Yes/No cell rather than on position has to survive that.
        page = "\n".join([
            "Listing of known creditors",
            "Name", "Address", "Related Party", "Creditor Type", "Amount",
            "Acme Building Supplies Pty Ltd", "12 Industry Rd Dandenong VIC",
            "No", "42,500.00",
        ])
        rows = creditor_tables.parse_cells(
            page.splitlines(), "Bust Co", "m1", "worrells")
        assert len(rows) == 1
        assert rows[0].creditor_name == "Acme Building Supplies Pty Ltd"
        assert rows[0].amount_aud == 42500.0


class TestLiveTwoAmountListing:
    """The two-amount layout, verbatim from two live Initial Advice reports.

    The September intake reported 80 creditors at exactly $0.00, and four
    documents reported $0.00 for every creditor. The cause is in these pages:
    the listing carries two amount columns, "ROCAP Amount" (from the
    director's report on company affairs) and "Identified Amount" (the
    liquidator's own figure). The ROCAP column reads $0.00 on almost every
    row, and the parser was taking the first amount cell after the Related
    Party flag - so the Australian Taxation Office came through at $0.00 when
    the document said $370,993.81.

    Both pages are captured by the diagnose workflow, cell for cell.
    """

    # FRONTLINE TRADES SOLUTIONS PTY LTD, page 17. Every row carries both
    # amount columns.
    FRONTLINE = "\n".join([
        "17",
        "D.",
        "Listing of known creditors (identifying related parties)",
        "Name", "Address", "Related Party", "ROCAP Amount", "Identified",
        "Amount",
        "Australian Taxation Office",
        "Australian Taxation Office, PO Box 9003",
        "NSW 2740",
        "No", "$0.00", "$370,993.81",
        "Commissioner of State Revenue",
        "(Queensland Revenue Office)",
        "PO Box 15931, City East    QLD 4002",
        "No", "$0.00", "$637.46",
        "New South Lawyers",
        "PO Box 1449, Parramatta   NSW 2124",
        "No", "$0.00", "$17,074.80",
        "PJH Lawyers",
        "Level 1, 60 Martin Place   NSW 2000",
        "No", "$0.00", "$1,138.50",
        "QUEENSLAND HOME WARRANTY",
        "SCHEME",
        "GPO Box 5099 BRISBANE QLD",
        "No", "$0.00", "$236.30",
        "WorkCover Queensland",
        "GPO Box 2772   QLD 4001",
        "No", "$0.00", "$14,638.01",
        "Lodge your claim online from the file's File Information page",
    ])

    # JC Mechanical Repairs Pty Ltd, page 30. Here most rows carry only one
    # amount cell - the Identified column is blank and PyMuPDF emits nothing
    # for it - so the row shape varies within a single table.
    JC_MECHANICAL = "\n".join([
        "30",
        "D.",
        "Listing of known creditors (identifying related parties)",
        "Name", "Address", "Related Party", "ROCAP Amount", "Identified",
        "Amount",
        "A.C.N. 603 303 126 PTY LTD",
        "Level 8, 360 Collins Street  Melbourne",
        "Victoria 3000",
        "No", "$0.00",
        "Australian Taxation Office (Insolvencies)",
        "No", "$0.00",
        "AZORA ASSET FINANCE PTY LTD",
        "PO Box 1915  Castle Hill NSW 1765",
        "No", "$0.00",
        "Commonwealth Bank",
        "Locked Bag 790  PARRAMATTA NSW",
        "2124",
        "No", "$0.00", "$18,872.25",
        "Fair Entitlements Guarantee",
        "No", "$0.00",
        "FLEXICOMMERCIAL PTY LTD",
        "LEVEL 1 , 121 Harrington Street  The",
        "Rocks NSW 2000",
        "No", "$0.00",
        "iCare Workers Insurance",
        "No", "$0.00",
        "IQumulate Premium Funding Pty Ltd",
    ])

    def frontline(self):
        return creditor_tables.parse_cells(
            self.FRONTLINE.splitlines(), "FRONTLINE TRADES SOLUTIONS PTY LTD",
            "m1", "worrells")

    def jc(self):
        return creditor_tables.parse_cells(
            self.JC_MECHANICAL.splitlines(), "JC Mechanical Repairs Pty Ltd",
            "m2", "worrells")

    def test_the_identified_amount_is_read_not_the_rocap_zero(self):
        amounts = {r.creditor_name: r.amount_aud for r in self.frontline()}
        assert amounts["Australian Taxation Office"] == 370993.81
        assert amounts["Commissioner of State Revenue"] == 637.46
        assert amounts["New South Lawyers"] == 17074.80
        assert amounts["PJH Lawyers"] == 1138.50
        assert amounts["WorkCover Queensland"] == 14638.01

    def test_every_frontline_creditor_has_a_stated_amount(self):
        rows = self.frontline()
        assert len(rows) == 6
        assert all(r.amount_known for r in rows)
        assert not any(r.amount_aud == 0.0 for r in rows)

    def test_the_second_amount_column_header_is_not_a_creditor(self):
        names = [r.creditor_name for r in self.frontline()]
        assert "Identified" not in names and "Amount" not in names

    def test_a_row_with_one_amount_cell_still_reads_it(self):
        amounts = {r.creditor_name: r.amount_aud for r in self.jc()}
        assert amounts["Commonwealth Bank"] == 18872.25

    def test_rows_with_only_a_zero_are_unquantified_not_zero(self):
        # A creditor owed nothing would not appear in a creditor listing. A
        # row whose every amount cell reads $0.00 is an unstated debt, and
        # recording it as $0.00 hands it to the $5,000 floor to be dropped as
        # "too small" - losing a genuine prospect for a reason that is false.
        rows = {r.creditor_name: r for r in self.jc()}
        assert not rows["AZORA ASSET FINANCE PTY LTD"].amount_known
        assert not rows["FLEXICOMMERCIAL PTY LTD"].amount_known
        assert rows["Commonwealth Bank"].amount_known

    def test_unquantified_rows_are_not_dropped_by_the_exposure_floor(self):
        prospects = qualify.apply(aggregate.build(self.jc()))
        flexi = [p for p in prospects if "FLEXICOMMERCIAL" in p.display_name.upper()]
        assert flexi and not any(
            p.disqualified_reason and "under the" in p.disqualified_reason
            for p in flexi
        )

    def test_all_seven_jc_creditors_are_read(self):
        assert [r.creditor_name for r in self.jc()] == [
            "A.C.N. 603 303 126 PTY LTD",
            "Australian Taxation Office (Insolvencies)",
            "AZORA ASSET FINANCE PTY LTD",
            "Commonwealth Bank",
            "Fair Entitlements Guarantee",
            "FLEXICOMMERCIAL PTY LTD",
            "iCare Workers Insurance",
        ]

    def test_a_wrapped_postcode_is_not_stripped_as_page_furniture(self):
        # "2124" arrives as its own numeric cell, indistinguishable from the
        # page number except by position: furniture sits above the first row.
        address = {r.creditor_name: r.address for r in self.jc()}
        assert address["Commonwealth Bank"] == \
            "Locked Bag 790  PARRAMATTA NSW 2124"

    def test_the_page_number_is_still_kept_out_of_the_first_row(self):
        names = [r.creditor_name for r in self.jc()]
        assert "30" not in names and "D." not in names
        assert names[0] == "A.C.N. 603 303 126 PTY LTD"


class TestLiveRfcvicListing:
    """The real creditor listing, end to end.

    Every string below is verbatim from page 18 of RFCVIC PTY LTD's published
    Initial Advice (trading as Reel Food Catering), captured by the diagnose
    workflow. This is the fixture that matters: the earlier synthetic ones were
    written to match the parser's assumptions and passed while the parser was
    incapable of reading a real document.
    """

    PAGE = "\n".join([
        "18",
        "D.",
        "Listing of known creditors (identifying related parties)",
        "Name", "Address", "Related Party", "ROCAP Amount",
        "Australian Alliance Automotive Finance Pty Limited",
        "Locked Bag 900  Milson Point NSW 1565", "No", "TBC",
        "Bidfood Australia Limited",
        "PO Box 220  Pendle Hill NSW 2145", "No", "TBC",
        "Silver Chef Rentals Pty Ltd",
        "PO Box 1760  Milton BC QLD 4064", "No", "TBC",
        "Velociti Capital Spv 1 Pty Ltd",
        "Unit 2, 4 ORRONG CRES  CAULFIELD NORTH VIC 3161", "No", "TBC",
        "Lodge your claim online from the file's File Information page",
    ])

    # The proposal response form carries Yes/No cells too, and must not be
    # mistaken for a listing.
    RESPONSE_FORM = "\n".join([
        "32", "H.", "Proposal response form and notices",
        "RFCVIC PTY LTD (In Liquidation)", "ACN: 678 250 327",
        "Yes", "No", "Object*",
        "Proposal 1 - Past fee approval", "Name of creditor:", "Address:",
    ])

    def rows(self):
        return creditor_tables.parse_cells(
            self.PAGE.splitlines(), "RFCVIC PTY LTD", "m1", "worrells")

    def test_the_listing_page_qualifies_despite_tbc_amounts(self):
        # Requiring amount cells on the page rejected this real listing:
        # every ROCAP Amount reads TBC, so the page has no amount cells.
        assert creditor_tables.is_creditor_table(self.PAGE)

    def test_the_proposal_response_form_does_not_qualify(self):
        assert not creditor_tables.is_creditor_table(self.RESPONSE_FORM)

    def test_all_four_creditors_are_read(self):
        assert [r.creditor_name for r in self.rows()] == [
            "Australian Alliance Automotive Finance Pty Limited",
            "Bidfood Australia Limited",
            "Silver Chef Rentals Pty Ltd",
            "Velociti Capital Spv 1 Pty Ltd",
        ]

    def test_page_number_and_section_letter_do_not_become_the_first_creditor(self):
        # "18" and "D." sit in the row buffer ahead of the first row. Left in,
        # the first creditor was lost entirely.
        names = [r.creditor_name for r in self.rows()]
        assert "18" not in names and "D." not in names
        assert names[0] == "Australian Alliance Automotive Finance Pty Limited"

    def test_rocap_amount_header_is_not_a_creditor(self):
        assert "ROCAP Amount" not in [r.creditor_name for r in self.rows()]

    def test_tbc_is_recorded_as_unknown_not_as_zero(self):
        assert all(not r.amount_known for r in self.rows())
        assert all(r.amount_aud == 0.0 for r in self.rows())

    def test_addresses_pair_with_the_right_creditor(self):
        rows = {r.creditor_name: r.address for r in self.rows()}
        assert rows["Bidfood Australia Limited"] == "PO Box 220  Pendle Hill NSW 2145"
        assert rows["Silver Chef Rentals Pty Ltd"] == "PO Box 1760  Milton BC QLD 4064"

    def test_unquantified_creditors_survive_the_exposure_floor(self):
        # The $5,000 floor cannot be applied to an unstated amount. Applying
        # it anyway would drop the entire early Worrells intake.
        prospects = qualify.apply(aggregate.build(self.rows()))
        assert all(not p.exposure_known for p in prospects)
        assert not any(
            p.disqualified_reason and "under the" in p.disqualified_reason
            for p in prospects
        )

    def test_only_the_trade_supplier_survives_qualification(self):
        # Of the four, Bidfood is the food wholesaler that supplied a catering
        # company on credit. The other three are finance, equipment rental and
        # a capital vehicle - none insurable as trade credit.
        prospects = qualify.apply(aggregate.build(self.rows()))
        assert [p.display_name for p in prospects if p.qualified] == [
            "Bidfood Australia Limited"
        ]

    def test_an_unquantified_prospect_scores_on_repeat_exposure_alone(self):
        prospect = aggregate.build(self.rows())[0]
        assert qualify.score(prospect) == 0
        two_matters = aggregate.build(
            self.rows()
            + [Creditor("Bidfood Australia Limited", "Other Bust Pty Ltd", "m2",
                        0.0, amount_known=False)]
        )
        bidfood = next(p for p in two_matters if "Bidfood" in p.display_name)
        assert qualify.score(bidfood) > 0


class TestRepeatExposureOnLiveData:
    """Bidfood appears in two of the three live listings.

    RFCVIC PTY LTD (Reel Food Catering) and SJMFood Pty Ltd both list Bidfood
    Australia as a creditor, spelled differently. That is the repeat-exposure
    signal the whole ranking is built on, observed in real data: one food
    wholesaler carrying bad debts from two separate food-service collapses.
    """

    RFCVIC = ["Australian Alliance Automotive Finance Pty Limited",
              "Bidfood Australia Limited", "Silver Chef Rentals Pty Ltd",
              "Velociti Capital Spv 1 Pty Ltd"]
    SJMFOOD = ["A.C.N. 603 273 365 PTY LTD", "ALLIED RETAIL FINANCE PTY LTD",
               "BIDFOOD AUSTRALIA LIMITED", "BOB & PETE'S PTY LIMITED"]

    def prospects(self):
        rows = [
            Creditor(n, "RFCVIC PTY LTD", "m1", 0.0, amount_known=False)
            for n in self.RFCVIC
        ] + [
            Creditor(n, "SJMFood Pty Ltd", "m2", 0.0, amount_known=False)
            for n in self.SJMFOOD
        ]
        return qualify.apply(aggregate.build(rows))

    def test_differently_spelled_bidfood_becomes_one_prospect(self):
        bidfood = [p for p in self.prospects() if "bidfood" in p.name_key]
        assert len(bidfood) == 1
        assert bidfood[0].matter_count == 2

    def test_bidfood_outranks_the_single_matter_prospects(self):
        qualified = [p for p in self.prospects() if p.qualified]
        assert qualified[0].matter_count == 2
        assert "BIDFOOD" in qualified[0].display_name.upper()

    def test_both_debtor_companies_are_recorded_against_it(self):
        bidfood = next(p for p in self.prospects() if "bidfood" in p.name_key)
        assert {m["debtor_company"] for m in bidfood.matters} == {
            "RFCVIC PTY LTD", "SJMFood Pty Ltd"
        }

    def test_the_financiers_are_all_excluded(self):
        excluded = {p.display_name for p in self.prospects() if not p.qualified}
        assert "Australian Alliance Automotive Finance Pty Limited" in excluded
        assert "ALLIED RETAIL FINANCE PTY LTD" in excluded
        assert "Silver Chef Rentals Pty Ltd" in excluded
        assert "Velociti Capital Spv 1 Pty Ltd" in excluded


class TestWorrellsHarvest:
    """The leg that turns a Worrells matter into creditors with no purchase.

    Stubs the HTTP client, because the portal is not reachable from the test
    environment. The page and document content are the live shapes captured by
    the diagnose workflow.
    """

    DETAIL_HTML = """<html><body>
      File Information RFCVIC PTY LTD File Details
      ACN/Estate#: 678 250 327 Office Name: Melbourne Trading Name None known
      Principal R Crispino Admin Type Creditors Vol Manager: D Hayman
      Start Date 08/09/2026 Exec Analyst D Hayman Status Priority
      Contact Person D Hayman Industry Accommodation and Food Services
      Appointee Roberto Crispino I want to Lodge a Proof of Debt
      <a href="/WebDocuments/1/initial.pdf">Initial Advice</a>
      <a href="/WebDocuments/1/first.pdf">First Advice</a>
    </body></html>"""

    LISTING = "\n".join([
        "18", "D.", "Listing of known creditors (identifying related parties)",
        "Name", "Address", "Related Party", "ROCAP Amount",
        "Bidfood Australia Limited", "PO Box 220  Pendle Hill NSW 2145", "No", "TBC",
        "Silver Chef Rentals Pty Ltd", "PO Box 1760  Milton BC QLD 4064", "No", "TBC",
    ])

    @staticmethod
    def pdf_bytes(text):
        import pymupdf

        doc = pymupdf.open()
        doc.new_page().insert_text((40, 50), text, fontsize=8)
        blob = doc.tobytes()
        doc.close()
        return blob

    class StubClient:
        """Serves the detail page and the documents; records what was fetched."""

        def __init__(self, detail_html, documents):
            self.detail_html = detail_html
            self.documents = documents
            self.fetched = []

        def get(self, url, **_):
            self.fetched.append(url)
            if url.endswith(".pdf"):
                name = url.rsplit("/", 1)[-1]
                if name not in self.documents:
                    raise RuntimeError(f"404 {name}")
                return type("R", (), {"content": self.documents[name]})()
            return type("R", (), {"text": self.detail_html})()

    def record(self):
        return {
            "matter_id": "m1",
            "company_name": "RFCVIC PTY LTD",
            "source": "worrells",
            "source_id": "202502200334422382761DFEC99D43BE",
        }

    def test_a_listing_is_harvested_into_creditors(self):
        client = self.StubClient(
            self.DETAIL_HTML, {"initial.pdf": self.pdf_bytes(self.LISTING)})
        rows, status, _ = worrells.harvest(self.record(), client)
        assert status == "ok"
        assert [r.creditor_name for r in rows] == [
            "Bidfood Australia Limited", "Silver Chef Rentals Pty Ltd"]

    def test_the_detail_panel_supplies_the_acn(self):
        # The ACN is what lets a Worrells matter reconcile against the same
        # company in the ASIC workbook.
        client = self.StubClient(
            self.DETAIL_HTML, {"initial.pdf": self.pdf_bytes(self.LISTING)})
        _, _, updates = worrells.harvest(self.record(), client)
        assert updates["acn"] == "678250327"
        assert updates["industry"] == "Accommodation and Food Services"
        assert updates["practitioner_firm"] == "Worrells"

    def test_the_best_ranked_document_wins_and_the_rest_are_not_fetched(self):
        # Initial Advice carries the listing 4 times out of 4; paying to
        # download the First Advice as well is wasted portal load.
        client = self.StubClient(self.DETAIL_HTML, {
            "initial.pdf": self.pdf_bytes(self.LISTING),
            "first.pdf": self.pdf_bytes(self.LISTING),
        })
        worrells.harvest(self.record(), client)
        assert any("initial.pdf" in u for u in client.fetched)
        assert not any("first.pdf" in u for u in client.fetched)

    def test_a_matter_with_no_documents_is_not_an_error(self):
        # Most of the New Appointments list is days old and carries nothing.
        client = self.StubClient(
            "<html><body>File Details ACN/Estate#: 123 456 789</body></html>", {})
        rows, status, _ = worrells.harvest(self.record(), client)
        assert (rows, status) == ([], "no-documents")

    def test_documents_without_a_listing_report_no_section(self):
        client = self.StubClient(self.DETAIL_HTML, {
            "initial.pdf": self.pdf_bytes("Remuneration report\nFees: 3,405.70"),
            "first.pdf": self.pdf_bytes("Covering letter only"),
        })
        rows, status, _ = worrells.harvest(self.record(), client)
        assert (rows, status) == ([], "no-section")

    def test_a_portal_failure_is_reported_not_raised(self):
        class Dead:
            def get(self, url, **_):
                raise RuntimeError("connection reset")

        rows, status, updates = worrells.harvest(self.record(), Dead())
        assert (rows, status, updates) == ([], "failed", {})

    def test_the_creditors_carry_the_document_they_came_from(self):
        client = self.StubClient(
            self.DETAIL_HTML, {"initial.pdf": self.pdf_bytes(self.LISTING)})
        rows, _, _ = worrells.harvest(self.record(), client)
        assert all(r.source_document == "Initial Advice" for r in rows)
        assert all(r.source == "worrells" for r in rows)
    def test_harvested_creditors_carry_the_debtor_industry(self, tmp_path, monkeypatch):
        # The industry is the column that tells a rep a timber supplier's bad
        # debts all came from construction. The Worrells path lost it once.
        from creditor_sourcing import cli

        client = self.StubClient(
            self.DETAIL_HTML, {"initial.pdf": self.pdf_bytes(self.LISTING)})
        monkeypatch.setattr(cli, "Client", lambda *a, **k: client)
        monkeypatch.setattr(worrells, "Client", lambda *a, **k: client)
        monkeypatch.setattr(ledger, "MATTERS", tmp_path / "matters.json")
        monkeypatch.setattr(ledger, "QUEUE", tmp_path / "queue.json")

        record = self.record()
        record["first_seen"] = date.today().isoformat()
        ledger.save_matters({record["matter_id"]: record})

        args = cli.build_parser().parse_args(["watch"])
        args.sources = ["worrells"]
        args.out = str(tmp_path / "creditors.json")
        cli.cmd_watch(args)

        import json

        rows = json.loads((tmp_path / "creditors.json").read_text())
        assert rows, "watch produced no creditors"
        assert all(
            r["debtor_industry"] == "Accommodation and Food Services" for r in rows
        )

    def test_re_harvesting_a_matter_replaces_its_creditors(self, tmp_path):
        # A practitioner lodging a fuller document should correct the data,
        # not duplicate every creditor already recorded against that matter.
        from creditor_sourcing.cli import _append_creditors

        path = tmp_path / "creditors.json"
        first = [Creditor("Old Name Pty Ltd", "RFCVIC PTY LTD", "m1", 0.0)]
        second = [Creditor("Bidfood Australia Limited", "RFCVIC PTY LTD", "m1", 0.0)]
        other = [Creditor("Untouched Pty Ltd", "Other Co", "m2", 0.0)]

        _append_creditors(path, first + other)
        _append_creditors(path, second)

        import json

        names = {r["creditor_name"] for r in json.loads(path.read_text())}
        assert names == {"Bidfood Australia Limited", "Untouched Pty Ltd"}

    def test_a_reparse_that_yields_nothing_clears_the_old_rows(self, tmp_path):
        # The parser fix that stopped producing "Report for NAVIQ GROUP PTY
        # LTD (Administrator Appointed)" left the row in place: the matter
        # re-parsed to zero creditors, and only matters present in the new
        # batch were being replaced. The fix landed and the bad creditor
        # stayed in the workbook.
        from creditor_sourcing.cli import _append_creditors

        path = tmp_path / "creditors.json"
        stale = [Creditor("Report for NAVIQ GROUP PTY LTD (Administrator "
                          "Appointed)", "NAVIQ GROUP PTY LTD", "m1", 747812.53)]
        other = [Creditor("Untouched Pty Ltd", "Other Co", "m2", 0.0)]
        _append_creditors(path, stale + other)

        _append_creditors(path, [], reparsed={"m1"})

        import json

        names = {r["creditor_name"] for r in json.loads(path.read_text())}
        assert names == {"Untouched Pty Ltd"}

    def test_a_matter_that_was_not_read_keeps_its_rows(self, tmp_path):
        # A transport failure or a scanned document is not evidence that the
        # creditors already recorded are wrong.
        from creditor_sourcing.cli import _append_creditors

        path = tmp_path / "creditors.json"
        held = [Creditor("Bidfood Australia Limited", "RFCVIC PTY LTD", "m1", 0.0)]
        _append_creditors(path, held)
        _append_creditors(path, [], reparsed=set())

        import json

        assert [r["creditor_name"] for r in json.loads(path.read_text())] == [
            "Bidfood Australia Limited"]



class TestOpenMatterLifecycle:
    """Which matters a run re-checks, and which it stops chasing."""

    @staticmethod
    def matter(**kw):
        base = {"matter_id": "m1", "company_name": "X", "source": "worrells",
                "first_seen": date.today().isoformat()}
        base.update(kw)
        return {base["matter_id"]: base}

    def test_a_matter_with_nothing_lodged_stays_open(self):
        # This is most of the Worrells intake and is exactly what we wait on.
        assert ledger.open_matters(self.matter(document_status="no-documents"))

    def test_a_captured_matter_closes(self):
        assert not ledger.open_matters(self.matter(creditors_captured=True))

    def test_documents_without_a_listing_stop_being_re_fetched(self):
        # Those documents will not grow a listing; re-fetching them weekly is
        # pure portal load for no possible gain.
        assert not ledger.open_matters(self.matter(document_status="no-section"))

    def test_a_scanned_document_stops_being_re_fetched(self):
        assert not ledger.open_matters(self.matter(document_status="scanned"))

    def test_a_transient_failure_leaves_the_matter_open(self):
        assert ledger.open_matters(self.matter(document_status="failed"))


class TestCliArgumentPositions:
    """-v must work before and after the subcommand.

    argparse puts top-level flags before the subcommand. That is easy to get
    wrong in a shell script, and it fails the entire run with "unrecognized
    arguments" rather than merely ignoring the flag. It silently broke the
    weekly-sourcing and probe workflows before either had ever been run.
    """

    @staticmethod
    def parse(argv):
        from creditor_sourcing.cli import build_parser

        return build_parser().parse_args(argv)

    @staticmethod
    def verbose(args):
        return getattr(args, "verbose", False)

    def test_verbose_before_the_subcommand(self):
        # The subparser must not write its own default over this.
        assert self.verbose(self.parse(["-v", "collect"]))

    def test_verbose_after_the_subcommand(self):
        assert self.verbose(self.parse(["collect", "--verbose"]))

    def test_quiet_when_the_flag_is_absent(self):
        assert not self.verbose(self.parse(["collect"]))

    def test_verbose_after_the_subcommand_with_other_flags(self):
        args = self.parse(["run", "--respect-schedule", "--verbose"])
        assert self.verbose(args) and args.respect_schedule

    @pytest.mark.parametrize(
        "stage", ["collect", "watch", "ingest", "report", "run", "schema", "probe"],
    )
    def test_every_stage_accepts_verbose_after_it(self, stage):
        argv = [stage, "asic-notices", "--verbose"] if stage == "probe" \
            else [stage, "--verbose"]
        assert self.verbose(self.parse(argv))

    def test_the_stage_is_still_required(self):
        with pytest.raises(SystemExit):
            self.parse(["--verbose"])


class TestAsicCheckOrder:
    """How a capped run spends its ASIC Connect requests.

    The 60-day window tracks ~1,600 open matters and grows every week. An
    uncapped pass is 1,600 requests at 400ms — eleven minutes that eventually
    outlasts the job timeout, most of it spent on appointment types that never
    produce a Form 5604.
    """

    @staticmethod
    def record(name, kind, appointed, checked=None):
        return {"company_name": name, "appointment_type": kind,
                "appointment_date": appointed, "last_checked": checked}

    def ordered(self, records):
        return [r["company_name"] for r in asic_connect.check_order(records)]

    def test_cvls_come_before_other_appointment_types(self):
        # Only a creditors' voluntary liquidation reliably produces a 5604.
        records = [
            self.record("court", "Court liquidation", "2026-09-05"),
            self.record("cvl", "Creditors' voluntary liquidation", "2026-09-01"),
        ]
        assert self.ordered(records)[0] == "cvl"

    def test_never_checked_comes_before_already_checked(self):
        # Otherwise a capped run re-asks about the same matters every week and
        # never reaches the backlog behind them.
        records = [
            self.record("checked", "Creditors' voluntary liquidation",
                        "2026-09-05", "2026-09-11"),
            self.record("fresh", "Creditors' voluntary liquidation", "2026-09-01"),
        ]
        assert self.ordered(records)[0] == "fresh"

    def test_newest_appointment_first_within_a_tier(self):
        # A recent CVL is the most likely to have just lodged its 5604.
        records = [
            self.record("older", "Creditors' voluntary liquidation", "2026-07-15"),
            self.record("newer", "Creditors' voluntary liquidation", "2026-09-01"),
        ]
        assert self.ordered(records) == ["newer", "older"]

    def test_the_least_recently_checked_rotates_to_the_front(self):
        records = [
            self.record("yesterday", "Creditors' voluntary liquidation",
                        "2026-09-05", "2026-09-11"),
            self.record("last month", "Creditors' voluntary liquidation",
                        "2026-09-05", "2026-08-11"),
        ]
        assert self.ordered(records)[0] == "last month"

    def test_nothing_is_dropped_by_ordering(self):
        records = [
            self.record(f"c{i}", "Court liquidation", "2026-09-05") for i in range(5)
        ]
        assert len(asic_connect.check_order(records)) == 5

    def test_missing_fields_do_not_raise(self):
        # Matters collected before a field existed, or with a blank type.
        assert len(asic_connect.check_order([{"company_name": "bare"}])) == 1


class TestAbnLookup:
    """Resolving the ABN that IRIS is searched by.

    Both title strings below are verbatim from ABN Lookup, captured by the
    diagnose workflow against real ACNs out of committed state.
    """

    SUELL = (
        "<html><head><title>Current details for ABN 50 683 236 259 | ABN Lookup"
        "</title></head><body>"
        "<table><tr><th>ABN</th><td>Active from 19 Dec 2024</td></tr>"
        "<tr><th>Entity name</th><td>SUELL EARTHMOVING PTY LTD</td></tr></table>"
        "</body></html>"
    )
    BREADROLL = (
        "<html><head><title>Current details for ABN 93 626 084 008 | ABN Lookup"
        "</title></head><body>BREADROLL ENTERPRISES PTY LTD</body></html>"
    )
    NOT_FOUND = (
        "<html><head><title>ACN not found | ABN Lookup</title></head>"
        "<body>No matching records</body></html>"
    )

    def test_the_abn_comes_from_the_title(self):
        assert abn_lookup.find_abn(self.SUELL) == "50683236259"
        assert abn_lookup.find_abn(self.BREADROLL) == "93626084008"

    def test_the_body_table_is_not_trusted_over_the_title(self):
        # Measured live: the cell after the "ABN" header holds the status,
        # "Active from 19 Dec 2024", not the number. Reading the table first
        # would return a date.
        assert abn_lookup.find_abn(self.SUELL) == "50683236259"

    def test_a_page_with_no_abn_yields_nothing(self):
        assert abn_lookup.find_abn(self.NOT_FOUND) is None

    @pytest.mark.parametrize(
        "abn", ["50 683 236 259", "93626084008", "51824753556"],
    )
    def test_real_abns_pass_the_checksum(self, abn):
        assert abn_lookup.is_valid_abn(abn)

    @pytest.mark.parametrize(
        # Eleven digits that are not an ABN: a transposition of a real one, a
        # run of zeros, and too few digits. The checksum is what stops a phone
        # number or a registration number being handed to a rep as an ABN.
        "text", ["50 683 236 295", "00000000000", "683236259", "", "abc"],
    )
    def test_non_abns_are_refused(self, text):
        assert not abn_lookup.is_valid_abn(text)

    def test_formatting_matches_how_abn_lookup_prints_it(self):
        assert abn_lookup.format_abn("50683236259") == "50 683 236 259"

    def test_the_search_url_is_the_one_that_works(self):
        # /ABN/View?abn=<acn> returns "ACN not found"; the search path
        # redirects to the record. Measured live on both.
        url = abn_lookup.search_url("683236259")
        assert url.startswith("https://abr.business.gov.au/Search/ResultsActive")
        assert "683236259" in url


class TestAFailedLookupIsNotAnAnswer:
    """A matter whose lookup failed must not be recorded as checked.

    The first run reported 744 creditors' voluntary liquidations checked with
    zero Form 5604s found. Every one of those requests had returned HTTP 404 -
    the endpoint does not serve that page - and the failure was swallowed per
    company, so "no 5604 exists" and "we never got an answer" were recorded
    identically. A whole leg of the pipeline looked healthy while doing
    nothing.
    """

    class Boom:
        def get(self, url, **kwargs):
            raise RuntimeError("404 Client Error: Not Found")

    class Empty:
        def get(self, url, **kwargs):
            class Response:
                text = "<html><body>No documents</body></html>"
            return Response()

    def matter(self):
        return Matter(source="asic", company_name="SUELL EARTHMOVING PTY LTD",
                      acn="683236259")

    def test_an_unreachable_lookup_is_reported_as_unreached(self):
        matter, reached = asic_connect.check(self.matter(), self.Boom())
        assert not reached
        assert matter.last_checked is None
        assert not matter.form_5604_lodged

    def test_a_real_answer_with_no_5604_is_reported_as_reached(self):
        matter, reached = asic_connect.check(self.matter(), self.Empty())
        assert reached
        assert matter.last_checked
        assert not matter.form_5604_lodged

    def test_a_matter_with_no_acn_is_not_retried_forever(self):
        bare = Matter(source="asic", company_name="No ACN Pty Ltd", acn=None)
        matter, reached = asic_connect.check(bare, self.Boom())
        assert reached and matter.last_checked


class TestIndividualsAreNotProspects:
    """A person is not a business with a receivables ledger.

    Individuals were 24 of the 142 qualified prospects in the 12 September run,
    four of them over $300,000. The documents give no signal - one creditor row
    in 452 carried the "Withheld due to privacy legislation" address - so the
    only evidence is the name, and the rule is one-sided: a name is a person
    only when it carries no business signal at all.

    Every name below is verbatim from that run.
    """

    @pytest.mark.parametrize(
        "name",
        ["Nadia Montesano", "Chad Aaron Gardiner", "Maisie Sansovini",
         "Garry Foreman", "Karen McIntosh", "Blagoja Stojanovski",
         "Li Xiang Feng", "Weihua Li", "Ajay Kumar", "Matthew Job",
         "Cal Bail", "John Pye",
         # Two people on one row, with the care-of marker that bled in from
         # the address column.
         "Barry Daniels and Laura Daniels C/-",
         "Jennifer Ann Taylor and Alan Taylor C/-"],
    )
    def test_people_are_recognised(self, name):
        assert qualify.looks_like_a_person(name)

    @pytest.mark.parametrize(
        "name",
        # Real businesses from the same run. The first six read like personal
        # names to a careless rule: two capitalised words and nothing else.
        ["The Meat Place", "Trade Risk", "Best Build Melb", "Studworks",
         "Bowens", "Masterlite Window Installations",
         "Gross Waddell",  # a commercial property agency, in individuals.yml
         "Dahlsens Building Centres", "KGF Cleaning Team", "Ability Plaster",
         "MBS Architectural", "Archiclad Pty Ltd", "Mitre 10",
         "Caves Beach Holdings Pty Ltd C/-", "PCG Development 3",
         "Century & Co", "Bidfood Australia Limited"],
    )
    def test_businesses_are_not_mistaken_for_people(self, name):
        assert not qualify.looks_like_a_person(name)

    def test_an_individual_is_dropped_with_a_reason(self):
        rows = [
            Creditor("Nadia Montesano", "MULTITUDE PLASTER PTY LTD", "m1", 711946.0),
            Creditor("Ability Plaster", "MULTITUDE PLASTER PTY LTD", "m1", 307715.0),
        ]
        prospects = qualify.apply(aggregate.build(rows))
        by_name = {p.display_name: p for p in prospects}
        assert not by_name["Nadia Montesano"].qualified
        assert "Individual" in by_name["Nadia Montesano"].disqualified_reason
        assert by_name["Ability Plaster"].qualified

    def test_a_title_is_enough_on_its_own(self):
        assert qualify.looks_like_a_person("Mr J Smith")
        assert qualify.looks_like_a_person("Dr Helen Nguyen")


class TestCareOfMarkers:
    """"C/-" at the end of a name is address text, not part of the name."""

    def test_the_marker_is_trimmed_off_the_name(self):
        page = "\n".join([
            "Listing of known creditors",
            "Name", "Address", "Related Party", "Amount",
            "Caves Beach Holdings Pty Ltd C/-",
            "Sutton Laurence King Lawyers Level 3, 405 Collins Street", "No",
            "141,216.00",
            "Barry Daniels and Laura Daniels C/-",
            "Level 3, 405 Collins Street Melbourne VIC 3000", "No",
            "203,579.00",
        ])
        names = [
            r.creditor_name for r in creditor_tables.parse_cells(
                page.splitlines(), "Some Debtor Pty Ltd", "m1", "worrells")
        ]
        assert names == [
            "Caves Beach Holdings Pty Ltd", "Barry Daniels and Laura Daniels"]


class TestNonTradeCreditorsFromTheLiveList:
    """Names that parsed correctly but are not trade credit prospects.

    All were qualified prospects in the 12 September run, several at the top
    of the list by score. A trade credit policy insures a supplier's
    receivables ledger, so the prospect has to be a business that sold goods
    or services on credit terms.
    """

    @pytest.mark.parametrize(
        ("name", "category"),
        [
            # Toll accounts: a recurring service billed on account, same shape
            # as a utility bill. Linkt was the top-scoring prospect on four
            # appearances.
            ("Linkt", "landlords_utilities"),
            ("Transurban Limited", "landlords_utilities"),
            # Card issuers: revolving credit, not a receivables ledger.
            ("AMEX", "financiers"),
            ("American Express Australia Limited", "financiers"),
            # Non-bank lender.
            ("Bizcap Au Pty Ltd", "financiers"),
            # Trustee vehicles lending into the company.
            ("DOH Investment Pty Ltd ATF the DOH Investment Trust",
             "related_party"),
            ("Smith Holdings as trustee for the Smith Family Trust",
             "related_party"),
            # No trading name to call.
            ("A.C.N. 603 303 126 PTY LTD", "noise"),
            ("ACN 002 738 472 Pty Ltd", "noise"),
        ],
    )
    def test_excluded_with_the_right_reason(self, name, category):
        hit = qualify.non_trade_reason(name)
        assert hit is not None, f"{name} was not excluded"
        assert hit[0] == category

    @pytest.mark.parametrize(
        "name",
        # Real trade suppliers from the same run. The patterns above must not
        # reach them: "Ability Plaster" contains no financier word, and
        # "Dahlsens Building Centres" is a timber and hardware merchant.
        ["Ability Plaster", "Dahlsens Building Centres", "Archiclad Pty Ltd",
         "YE Commercial Interiors", "Studworks", "MBS Architectural",
         "Melbourne Plaster Labour services", "Bidfood Australia Limited"],
    )
    def test_real_suppliers_are_not_excluded(self, name):
        assert qualify.non_trade_reason(name) is None


class TestArtefactsFromTheSecondLiveRun:
    """Three names that reached the qualified prospect list but are not companies.

    All from the 12 September run over 31 real insolvencies. Each is a
    different failure: a running page header, a ledger column header, and an
    employee entitlement category. All three would have been handed to the
    sales team as businesses to call.
    """

    def test_the_running_page_header_is_not_a_creditor(self):
        # "Report for NAVIQ GROUP PTY LTD (Administrator Appointed)" is the
        # header printed on every page of the report. It reached the list
        # owed $747,812.53 - a company cannot be its own creditor.
        page = "\n".join([
            "Listing of known creditors",
            "Name", "Address", "Related Party", "Amount",
            "Report for NAVIQ GROUP PTY LTD (Administrator Appointed)",
            "6. Unsecured Creditors A list of the known unsecured creditors",
            "No", "747,812.53",
            "Archiclad Pty Ltd", "5 Foundry Rd Sunshine VIC 3020", "No",
            "484,312.00",
        ])
        rows = creditor_tables.parse_cells(
            page.splitlines(), "NAVIQ GROUP PTY LTD", "m1", "worrells")
        assert [r.creditor_name for r in rows] == ["Archiclad Pty Ltd"]

    def test_the_line_parser_also_rejects_the_debtor_itself(self):
        rows = creditor_tables.parse_lines(
            ["Report for TOWM FOOD PTY LTD (In Liquidation) 98,533.00",
             "Bidfood Australia Limited 12,400.00"],
            "TOWM FOOD PTY LTD", "m1", "worrells")
        assert [r.creditor_name for r in rows] == ["Bidfood Australia Limited"]

    def test_a_ledger_column_header_is_not_a_creditor(self):
        # "Debit Amount" became a creditor owed $984,910 and demoted the real
        # creditor, the ATO, to its address.
        assert creditor_tables.is_header_fragment("Debit Amount")
        page = "\n".join([
            "Listing of known creditors",
            "Name", "Address", "Related Party", "Debit Amount",
            "Debit Amount",
            "Australian Taxation Office (Insolvencies)",
            "No", "984,910.00",
        ])
        rows = creditor_tables.parse_cells(
            page.splitlines(), "MAGNATE INTERNATIONAL PTY LTD", "m1", "worrells")
        assert [r.creditor_name for r in rows] == [
            "Australian Taxation Office (Insolvencies)"]

    def test_an_entitlement_category_is_not_a_prospect(self):
        # "Annual Leave", address "N/A", owed $27,798.47 on HOMEMAKERS S.C's
        # listing. A real row in the document, but not a business to call.
        rows = [
            Creditor("Annual Leave", "HOMEMAKERS S.C PTY LTD", "m1", 27798.47),
            Creditor("Bidfood Australia Limited", "HOMEMAKERS S.C PTY LTD",
                     "m1", 27798.47),
        ]
        prospects = qualify.apply(aggregate.build(rows))
        assert [p.display_name for p in prospects if p.qualified] == [
            "Bidfood Australia Limited"]


class TestHeaderFragments:
    """Column headers must never become creditors.

    Every string here came out of the first live pipeline run over 29 real
    insolvencies. A two-amount table ("ROCAP / Identified") splits its header
    across cells, so fragments reached the row buffer and became creditor
    names: ten rows across nine insolvencies, one carrying $160,000. A
    prospect named "Identified" owed $160k would have reached the sales team.
    """

    @pytest.mark.parametrize(
        "cell",
        ["ROCAP / Identified", "Identified", "ROCAP /", "Estimated Amount",
         "Related Party", "Amount Owed", "Creditor Name", "Total", "Yes", "No"],
    )
    def test_header_vocabulary_is_rejected(self, cell):
        assert creditor_tables.is_header_fragment(cell)

    @pytest.mark.parametrize(
        "name",
        # All real creditors from the live run. The first three start with
        # "No" and the rest contain a header word - a cruder rule would have
        # deleted genuine companies.
        ["No Splash Concrete Pumping", "NOBLE INSUL & CO Pty Ltd",
         "Noteg Pty Ltd", "Realtime Flowers", "Lynch Group (Flower HQ)",
         "Coast Cafe Supplies", "Petal Peddlers", "Square Australia Pty Ltd",
         "Identified Pty Ltd", "Total Tools Pty Ltd", "Balance Nutrition Pty Ltd"],
    )
    def test_real_companies_survive(self, name):
        assert not creditor_tables.is_header_fragment(name)

    def test_a_two_amount_table_does_not_leak_its_header(self):
        page = "\n".join([
            "Listing of known creditors",
            "Name", "Address", "Related Party", "ROCAP /", "Identified",
            "Coast Cafe Supplies", "12 Trade St Brisbane QLD", "No",
            "406.00", "406.00",
            "Realtime Flowers", "9 Market Rd Sydney NSW", "No",
            "110,100.18", "110,100.18",
        ])
        rows = creditor_tables.parse_cells(
            page.splitlines(), "Kor Enterprise Pty Ltd", "m1", "worrells")
        assert [r.creditor_name for r in rows] == [
            "Coast Cafe Supplies", "Realtime Flowers"]
        assert rows[0].amount_aud == 406.0


class TestPolicyListExcludesExistingClients:
    """Item 1 of the handover: the PolicyList had never been run.

    Against the 4 June 2026 export, eight of the 115 prospects on the 13
    September list were current NCI clients. Five matched on the client
    name; Mitre 10 and Home Timber & Hardware only appear in the POLICY name
    column (their client is TOTAL TOOLS & HARDWARE GROUP), and Studworks is
    how a creditor listing abbreviates STUDWORKS PROFILE SYSTEMS.
    """

    POLICIES = [
        ("NCI8854V", "MITRE 10 AUSTRALIA PTY LTD", "TOTAL TOOLS & HARDWARE GROUP"),
        ("NCI8852V", "HOME TIMBER & HARDWARE GROUP PTY LTD", "TOTAL TOOLS & HARDWARE GROUP"),
        ("NCI7912V", "STUDWORKS PROFILE SYSTEMS PTY LTD", "STUDWORKS PROFILE SYSTEMS PTY LTD"),
        ("NCI8737V", "SUPAPANEL AUSTRALIA PTY LTD", "SUPAPANEL AUSTRALIA PTY LTD"),
        ("NCI7792V", "FETCH PERSONNEL PTY LTD", "FETCH PERSONNEL PTY LTD"),
        ("898191", "AMERICOLD LOGISTICS LTD", "AMERICOLD LOGISTICS LTD"),
        ("NCI8206N", "ANTEC GROUP PTY LIMITED", "METAL MANUFACTURES PTY LIMITED"),
        ("NCI5366V", "ACROW FORMWORK AND SCAFFOLDING PTY LTD", "ACROW FORMWORK AND SCAFFOLDING PTY LTD"),
        ("NCI7801S", "CENTURY PRODUCTS (S.A.) PROPRIETARY LIMITED", "CENTURY PRODUCTS (S.A.) PROPRIETARY LIMITED"),
        ("AU25143500", "WELLPHARM PTY LTD AND MATTHEW BELLGROVE PHARMACY PTY LTD", "WELLPHARM PTY LTD"),
        ("NCI0001", "MELBOURNE COMMERCIAL CARPENTRY PTY LTD", "MELBOURNE COMMERCIAL CARPENTRY PTY LTD"),
        ("NCI0002", "MELBOURNE COMMERCIAL CLEANING PTY LTD", "MELBOURNE COMMERCIAL CLEANING PTY LTD"),
    ]

    @pytest.fixture
    def keys(self):
        from creditor_sourcing.enrich import policylist
        return policylist.fold_names(
            name for _, policy, client in self.POLICIES for name in (policy, client)
        )

    @pytest.mark.parametrize(
        ("creditor", "policyholder"),
        [
            ("Supapanel australia", "SUPAPANEL AUSTRALIA PTY LTD"),
            ("Fetch Personnel", "FETCH PERSONNEL PTY LTD"),
            ("AMERICOLD LOGISTICS LIMITED", "AMERICOLD LOGISTICS LTD"),
            ("METAL MANUFACTURES PTY LIMITED", "METAL MANUFACTURES PTY LIMITED"),
            ("Acrow Formwork And Scaffolding Pty Ltd",
             "ACROW FORMWORK AND SCAFFOLDING PTY LTD"),
            # Policy-name column only.
            ("Mitre 10", "MITRE 10 AUSTRALIA PTY LTD"),
            ("Home Timber & Hardware Group Pty Ltd",
             "HOME TIMBER & HARDWARE GROUP PTY LTD"),
            # Distinctive prefix of the policyholder's name.
            ("Studworks", "STUDWORKS PROFILE SYSTEMS PTY LTD"),
        ],
    )
    def test_existing_clients_match(self, keys, creditor, policyholder):
        assert qualify.policylist_match(creditor, keys) == policyholder

    @pytest.mark.parametrize(
        "creditor",
        [
            "Pharmacy",                      # one generic word, prefix rule refused
            "Century & Co",                  # shares a word with Century Products
            "Melbourne Plaster Labour services",
            "Melbourne Commercial",          # prefixes TWO policyholders - ambiguous
            "Dahlsens Building Centres",
            "",
        ],
    )
    def test_other_creditors_do_not_match(self, keys, creditor):
        assert qualify.policylist_match(creditor, keys) is None

    def test_apply_drops_a_client_with_a_visible_reason(self, keys):
        prospects = aggregate.build([creditor("Mitre 10", amount=51822.0)])
        [prospect] = qualify.apply(prospects, keys)
        assert prospect.qualified is False
        assert prospect.policylist_match == "MITRE 10 AUSTRALIA PTY LTD"
        assert "Existing NCI client" in prospect.disqualified_reason

    def _write_csv(self, path, header, rows, title_row=None):
        import csv
        with path.open("w", newline="", encoding="utf-8") as fh:
            writer = csv.writer(fh)
            if title_row is not None:
                writer.writerow(title_row)
            writer.writerow(header)
            writer.writerows(rows)

    def test_csv_loads_both_name_columns(self, tmp_path):
        from creditor_sourcing.enrich import policylist
        path = tmp_path / "policylist.csv"
        self._write_csv(
            path,
            ["policy_no", "policy_name", "client_name", "contact_first", "email", "state"],
            [[p, policy, client, "Kim", "kim@example.com", "VIC"]
             for p, policy, client in self.POLICIES],
        )
        keys = policylist.load(path)
        assert normalise_name("MITRE 10 AUSTRALIA PTY LTD") in keys
        assert normalise_name("TOTAL TOOLS & HARDWARE GROUP") in keys
        # Contact columns are never read as company names.
        assert "kim" not in keys

    def test_xlsx_export_with_a_title_row_above_the_header(self, tmp_path):
        from openpyxl import Workbook
        from creditor_sourcing.enrich import policylist
        path = tmp_path / "PolicyList.xlsx"
        wb = Workbook()
        ws = wb.active
        ws.append(["Policy List Report - 04/06/2026"])
        ws.append(["Policy No", "Policy Name", "Client Name", "State", "Industry"])
        for p, policy, client in self.POLICIES:
            ws.append([p, policy, client, "VIC", "BUILDING / HARDWARE"])
        wb.save(path)
        keys = policylist.load(path)
        assert len(keys) == len(policylist.fold_names(
            n for _, a, b in self.POLICIES for n in (a, b)))
        assert normalise_name("STUDWORKS PROFILE SYSTEMS PTY LTD") in keys

    def test_missing_or_headerless_file_switches_exclusion_off(self, tmp_path, caplog):
        from creditor_sourcing.enrich import policylist
        assert policylist.load(tmp_path / "nope.csv") == {}
        path = tmp_path / "junk.csv"
        self._write_csv(path, ["a", "b"], [["1", "2"]])
        with caplog.at_level("WARNING"):
            assert policylist.load(path) == {}
        assert "client exclusion is OFF" in caplog.text

    def test_the_committed_export_is_contact_free_and_loads(self):
        import csv
        from creditor_sourcing import config
        from creditor_sourcing.enrich import policylist
        path = config.REPO_ROOT / config.settings()["qualify"]["policylist_path"]
        with path.open(encoding="utf-8") as fh:
            header = next(csv.reader(fh))
        assert header == ["policy_no", "policy_name", "client_name", "state", "industry"]
        keys = policylist.load(path)
        assert len(keys) > 3000
        assert qualify.policylist_match("Mitre 10", keys) == "MITRE 10 AUSTRALIA PTY LTD"


class TestNonTradeCreditorsFromTheThirteenSeptemberList:
    """Twenty-five qualified prospects on the 13 September workbook that are
    not trade suppliers - fintech lenders, a fuel card, workers' compensation
    agents, an insurance broker, accountants, a utility written as one word,
    a building regulator, a body corporate and a bare category word."""

    @pytest.mark.parametrize(
        ("name", "category"),
        [
            ("Shift", "financiers"),
            ("Lumi", "financiers"),
            ("OnDeck", "financiers"),
            ("Dynamoney", "financiers"),
            ("MONEYME FINANCIAL GROUP PTY LTD", "financiers"),
            ("Square Australia Pty Ltd", "financiers"),
            ("Procuret Operating Pty Limited", "financiers"),
            ("Procuret Funding No. 5 Pty Ltd", "financiers"),
            ("FLEXICOMMERCIAL PTY LTD", "financiers"),
            ("NISSAN FINANCIAL SERVICES", "financiers"),
            ("Motorpass", "financiers"),
            ("Gallagher Bassett Services Workers Compensation VIC P/L", "insurance"),
            ("Gallagher Bassett Services", "insurance"),
            ("Trade Risk", "insurance"),
            ("KHI Partners", "professional_services"),
            ("DLK Advisory", "professional_services"),
            ("Platinum Advisory Accountants", "professional_services"),
            ("Bell Partners Newcastle", "professional_services"),
            ("Mergers and Acquisitions", "professional_services"),
            ("EnergyAustralia", "landlords_utilities"),
            ("The Owners - Units Plan No 4312", "landlords_utilities"),
            ("Building and Plumbing Commission", "statutory"),
            ("Pharmacy", "noise"),
        ],
    )
    def test_excluded_with_the_right_reason(self, name, category):
        hit = qualify.non_trade_reason(name)
        assert hit is not None, f"{name} was not excluded"
        assert hit[0] == category

    @pytest.mark.parametrize(
        "name",
        # Real suppliers from the same list that share a word or a shape with
        # the patterns above and must survive them.
        ["Bowens", "Crimsafe Security Systems", "Criterion Industries",
         "Nexdoor Systems", "Universal Fluid Power Pty Ltd", "Realtime Flowers",
         "The Meat Place", "Digital Horizons", "Night Shift Plastering Pty Ltd",
         "Square Peg Joinery", "Knauf Gypsum Pty Ltd", "Spicers Australia Pty Ltd",
         "Trumark Group", "Battmans Insulation Services", "RAR Developments"],
    )
    def test_real_suppliers_are_not_excluded(self, name):
        assert qualify.non_trade_reason(name) is None

    @pytest.mark.parametrize(
        "name", ["Phyllis (Meiping ) Yang", "William Longhurst 002"],
    )
    def test_decorated_person_names_are_people(self, name):
        assert qualify.looks_like_a_person(name) is True

    @pytest.mark.parametrize(
        "name", ["Mitre 10", "PCG Development 3", "Studio 54 Design", "M2 Plaster Pty Ltd"],
    )
    def test_a_digit_inside_a_name_still_means_business(self, name):
        assert qualify.looks_like_a_person(name) is False
