"""Build the weekly prospect workbook.

Five tabs, in the order a rep actually uses them:

  Prospects        qualified companies, highest priority first - the call list
  Repeat Exposure  companies bleeding across more than one administration
  Purchase Queue   the ASIC documents to buy - the one manual step
  Excluded         everything dropped, with the reason, so rules stay auditable
  Matters          every administration being watched and its 5604 status
"""

from __future__ import annotations

import logging
from collections.abc import Iterable, Sequence
from datetime import date
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from .models import Prospect

log = logging.getLogger(__name__)

HEADER_FILL = PatternFill("solid", fgColor="1F3864")
HEADER_FONT = Font(name="Arial", size=10, bold=True, color="FFFFFF")
BODY_FONT = Font(name="Arial", size=10)
MONEY = '"$"#,##0.00'

PRIORITY_FILL = PatternFill("solid", fgColor="C6EFCE")   # repeat exposure
CRM_FILL = PatternFill("solid", fgColor="FFF2CC")        # already in Pipedrive
DROPPED_FILL = PatternFill("solid", fgColor="F2F2F2")


def _sheet(workbook: Workbook, title: str, headers: Sequence[str]) -> Worksheet:
    sheet = workbook.create_sheet(title[:31])
    sheet.append(list(headers))
    for cell in sheet[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:{get_column_letter(len(headers))}1"
    sheet.row_dimensions[1].height = 28
    return sheet


def _finish(sheet: Worksheet, widths: Sequence[int]) -> None:
    for i, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(i)].width = width
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            if cell.font is BODY_FONT or cell.font.name != "Arial":
                cell.font = BODY_FONT


def _prospect_rows(sheet: Worksheet, prospects: Iterable[Prospect]) -> None:
    for prospect in prospects:
        debtors = "; ".join(
            f"{m['debtor_company']} "
            f"({'$' + format(m['amount_aud'], ',.0f') if m.get('amount_known', True) else 'TBC'})"
            for m in prospect.matters
        )
        sheet.append(
            [
                prospect.display_name,
                prospect.score,
                prospect.total_exposure_aud if prospect.exposure_known else "TBC",
                prospect.matter_count,
                debtors,
                ", ".join(prospect.debtor_industries),
                prospect.abn,
                prospect.state,
                prospect.contact_name,
                prospect.contact_title,
                prospect.contact_email,
                prospect.contact_phone,
                prospect.contact_source,
                prospect.pipedrive_org_name,
                prospect.pipedrive_owner,
                prospect.pipedrive_url,
            ]
        )
        row = sheet.max_row
        if prospect.exposure_known:
            sheet.cell(row=row, column=3).number_format = MONEY
        if prospect.matter_count > 1:
            for column in range(1, 6):
                sheet.cell(row=row, column=column).fill = PRIORITY_FILL
        if prospect.pipedrive_org_id:
            for column in range(14, 17):
                sheet.cell(row=row, column=column).fill = CRM_FILL


def build(
    prospects: list[Prospect],
    queue: list[dict[str, Any]],
    matters: dict[str, dict[str, Any]],
    out_dir: Path,
    prefix: str = "NCI_Creditor_Prospects",
) -> Path:
    workbook = Workbook()
    workbook.remove(workbook.active)

    qualified = [p for p in prospects if p.qualified]
    excluded = [p for p in prospects if not p.qualified]

    headers = [
        "Creditor (prospect)", "Score", "Total exposure", "# insolvencies",
        "Owed by (debtor companies)", "Debtor industries", "ABN", "State",
        "Contact", "Title", "Email", "Phone", "Contact source",
        "Pipedrive org", "Pipedrive owner", "Pipedrive link",
    ]
    widths = [38, 7, 16, 13, 52, 30, 15, 8, 24, 24, 30, 18, 14, 28, 20, 42]

    sheet = _sheet(workbook, "Prospects", headers)
    _prospect_rows(sheet, qualified)
    _finish(sheet, widths)

    repeat = _sheet(workbook, "Repeat Exposure", headers)
    _prospect_rows(repeat, [p for p in qualified if p.matter_count > 1])
    _finish(repeat, widths)

    queue_headers = [
        "Company (insolvent)", "ACN", "ABN (paste into IRIS)", "Form",
        "Document number", "Lodged", "Appointment type",
        "ASIC Connect link (buy here)", "IRIS debtor search", "Queued",
    ]
    queue_sheet = _sheet(workbook, "Purchase Queue", queue_headers)
    for row in queue:
        queue_sheet.append(
            [
                row.get("company_name"), row.get("acn"), row.get("abn"), "5604",
                row.get("document_number"), row.get("lodged_date"),
                row.get("appointment_type"), row.get("asic_connect_url"),
                row.get("iris_search_url"), row.get("queued_at"),
            ]
        )
    _finish(queue_sheet, [38, 14, 18, 8, 20, 13, 26, 62, 46, 20])

    excluded_headers = ["Creditor", "Total exposure", "# insolvencies",
                        "Why it was excluded"]
    excluded_sheet = _sheet(workbook, "Excluded", excluded_headers)
    for prospect in excluded:
        excluded_sheet.append(
            [
                prospect.display_name, prospect.total_exposure_aud,
                prospect.matter_count, prospect.disqualified_reason,
            ]
        )
        excluded_sheet.cell(row=excluded_sheet.max_row, column=2).number_format = MONEY
        for column in range(1, 5):
            excluded_sheet.cell(row=excluded_sheet.max_row, column=column).fill = DROPPED_FILL
    _finish(excluded_sheet, [38, 16, 13, 52])

    matter_headers = [
        "Insolvent company", "ACN", "Source", "Appointment type",
        "Appointment date", "Industry", "Industry (subdivision)", "State",
        "Postcode", "Practitioner", "Form 5604 lodged", "5604 date",
        "Document number", "Purchased", "Creditors captured", "Last checked",
    ]
    matter_sheet = _sheet(workbook, "Matters", matter_headers)
    for matter in sorted(matters.values(), key=lambda m: m.get("company_name") or ""):
        matter_sheet.append(
            [
                matter.get("company_name"), matter.get("acn"), matter.get("source"),
                matter.get("appointment_type"), matter.get("appointment_date"),
                matter.get("industry"), matter.get("industry_subdivision"),
                matter.get("state"), matter.get("postcode"),
                matter.get("practitioner"),
                "Yes" if matter.get("form_5604_lodged") else "No",
                matter.get("form_5604_date"), matter.get("form_5604_doc_number"),
                "Yes" if matter.get("form_5604_purchased") else "No",
                "Yes" if matter.get("creditors_captured") else "No",
                matter.get("last_checked"),
            ]
        )
    _finish(matter_sheet,
            [38, 14, 10, 26, 16, 22, 24, 16, 10, 26, 16, 13, 18, 11, 18, 14])

    out_dir.mkdir(parents=True, exist_ok=True)
    path = out_dir / f"{prefix}_{date.today().isoformat()}.xlsx"
    workbook.save(path)
    log.info(
        "Workbook: %s (%d prospects, %d repeat, %d excluded, %d to buy)",
        path.name, len(qualified),
        sum(1 for p in qualified if p.matter_count > 1), len(excluded), len(queue),
    )
    return path
