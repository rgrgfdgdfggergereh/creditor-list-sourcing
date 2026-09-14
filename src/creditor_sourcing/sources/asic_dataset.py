"""ASIC insolvency statistics workbook - the "data set" tab.

The workbook ASIC publishes at download.asic.gov.au carries summary tables for
Series 1 and Series 2, and also a row-per-appointment sheet named **"data set"**
listing the companies that entered external administration. That sheet is the
primary named-company feed for this pipeline: it is a structured file rather
than scraped markup, so it cannot break on a CSS change, and one download
replaces hundreds of page requests.

The published URL carries the publication date in its path and ASIC mints a new
media id each release, so `resolve_latest_url()` reads the insolvency statistics
landing page to find the current workbook rather than relying on a pinned URL
going stale.

Column handling
---------------
`HEADER_ALIASES` maps every spelling ASIC has used for a field onto the Matter
field it populates, and the header row is located by scanning for the first row
that matches at least two known headers - the sheet carries title and note rows
above the real header. If the company-name column cannot be found the loader
raises rather than returning an empty list, because "no insolvencies this week"
and "the schema moved" must never look the same.
"""

from __future__ import annotations

import io
import logging
import re
from datetime import date, datetime, timedelta
from pathlib import Path
from typing import Any

from bs4 import BeautifulSoup

from .. import config
from ..models import Matter
from .http import Client

log = logging.getLogger(__name__)

LANDING_PAGE = (
    "https://www.asic.gov.au/about-asic/corporate-publications/statistics/"
    "insolvency-statistics"
)
SHEET_CANDIDATES = ("data set", "dataset", "data_set", "raw data", "series 1 raw data")

# Members' voluntary liquidation is a solvent wind-up: there are no creditors
# taking a loss, so those rows are never prospect material.
SOLVENT_APPOINTMENTS = ("members' voluntary", "members voluntary")

# Lower-cased, punctuation-stripped header text -> Matter field.
#
# The exact spellings ASIC uses today are marked (actual). The rest are
# tolerated variants, kept so a rename does not break the run: ASIC has used
# several over the years and the sheet is republished monthly.
HEADER_ALIASES: dict[str, str] = {
    # Identity
    "organisation name": "company_name",          # actual
    "company name": "company_name",
    "entity name": "company_name",
    "name": "company_name",
    "company": "company_name",
    "acn no": "acn",                              # actual
    "acn": "acn",
    "a c n": "acn",
    "australian company number": "acn",
    # Appointment
    "appointment type": "appointment_type",       # actual
    "initial appointment type": "appointment_type",
    "type of appointment": "appointment_type",
    "external administration type": "appointment_type",
    "effective date": "appointment_date",         # actual
    "appointment date": "appointment_date",
    "date of appointment": "appointment_date",
    "date appointed": "appointment_date",
    "appointee person or company": "practitioner",  # actual
    "appointee": "practitioner",
    "practitioner": "practitioner",
    "practitioner name": "practitioner",
    "appointee name": "practitioner",
    "firm": "practitioner_firm",
    "practitioner firm": "practitioner_firm",
    # Industry - division is the headline, subdivision is the useful detail
    "industry type division": "industry",             # actual
    "industry type subdivision": "industry_subdivision",  # actual
    "industry": "industry",
    "industry division": "industry",
    "anzsic division": "industry",
    # Location. Principal place of business beats state of incorporation for
    # sales territory: it is where the company actually trades.
    "principal place of business state or territory": "state",   # actual
    "principal place of business postcode": "postcode",          # actual
    "state of incorporation state or territory": "state_of_incorporation",  # actual
    "state": "state",
    "state territory": "state",
    # Series 1 marks a company's FIRST appointment. Series 2 counts every
    # appointment including subsequent ones, which would duplicate matters.
    "series 1 companies entering": "series_1",    # actual
    "series 2 all appointments": "series_2",      # actual
}

_PUNCT = re.compile(r"[^a-z0-9 ]+")
_SPACE = re.compile(r"\s+")


def _norm_header(value: Any) -> str:
    """Fold a header cell to a lookup key.

    Headers carry embedded newlines ("Period\n(Year month)"), trailing spaces
    ("ACN No ") and parenthesised qualifiers, so punctuation and whitespace are
    flattened before lookup.
    """
    text = _PUNCT.sub(" ", str(value or "").strip().lower())
    return _SPACE.sub(" ", text).strip()


def _normalise_acn(value: Any) -> str | None:
    """Zero-pad an ACN back to nine digits.

    ACNs are stored as numbers in this sheet, so leading zeros are gone:
    TANCRED BROTHERS PTY LTD arrives as 25712 and is really 000 025 712.
    Dropping short values instead of padding them silently loses every company
    registered early enough to have a low ACN.
    """
    digits = re.sub(r"[^0-9]", "", str(value or ""))
    if not digits or len(digits) > 9:
        return None
    return digits.zfill(9)


def _is_true(value: Any) -> bool:
    """Read a Series 1 / Series 2 flag cell, which ASIC writes as 1 or blank."""
    if value in (None, ""):
        return False
    text = str(value).strip().lower()
    return text in {"1", "1.0", "y", "yes", "true"}


def resolve_latest_url(client: Client | None = None) -> str:
    """Find the current Series 1 and 2 workbook link on ASIC's landing page.

    Falls back to the configured URL if the page cannot be read, so a site
    redesign degrades to "last known good" rather than failing the run.
    """
    configured = config.settings()["sources"]["asic_stats"]["url"]
    client = client or Client()
    try:
        html = client.get(LANDING_PAGE).text
    except Exception as exc:  # noqa: BLE001
        log.warning("Could not read the ASIC statistics page (%s) - using %s", exc, configured)
        return configured

    soup = BeautifulSoup(html, "lxml")
    best: tuple[str, str] | None = None
    for link in soup.find_all("a", href=True):
        href = link["href"]
        if not href.lower().endswith(".xlsx"):
            continue
        slug = href.lower()
        if "insolvency-statistics" not in slug:
            continue
        # Prefer the combined Series 1 and Series 2 workbook - it is the one
        # that carries the data set sheet.
        if "series-1-and-series-2" in slug:
            return href if href.startswith("http") else "https://www.asic.gov.au" + href
        if best is None:
            best = (href, link.get_text(strip=True))

    if best:
        href = best[0]
        return href if href.startswith("http") else "https://www.asic.gov.au" + href
    log.warning("No workbook link found on the ASIC statistics page - using %s", configured)
    return configured


def download(url: str | None = None, client: Client | None = None) -> bytes:
    client = client or Client()
    url = url or resolve_latest_url(client)
    log.info("Downloading %s", url)
    return client.get(url).content


def _find_sheet(workbook) -> Any:
    names = {name.lower().strip(): name for name in workbook.sheetnames}
    for candidate in SHEET_CANDIDATES:
        if candidate in names:
            return workbook[names[candidate]]
    # Fall back to any sheet whose name contains "data".
    for lowered, original in names.items():
        if "data" in lowered:
            return workbook[original]
    raise RuntimeError(
        f"No data-set sheet in the workbook. Sheets present: {workbook.sheetnames}"
    )


def _find_header(rows: list[tuple], limit: int = 25) -> tuple[int, dict[int, str]]:
    """Locate the header row and map its columns onto Matter fields.

    The sheet has title and note rows above the header, so the header is found
    by content: the first row that matches at least two known header names.
    """
    for index, row in enumerate(rows[:limit]):
        mapping: dict[int, str] = {}
        for column, value in enumerate(row):
            field = HEADER_ALIASES.get(_norm_header(value))
            if field and field not in mapping.values():
                mapping[column] = field
        if len(mapping) >= 2 and "company_name" in mapping.values():
            return index, mapping
    raise RuntimeError(
        "Could not find a header row with a company-name column in the data set "
        "sheet. Run `python -m creditor_sourcing schema` and add the real "
        "headers to HEADER_ALIASES in sources/asic_dataset.py."
    )


def _parse_date(value: Any) -> str | None:
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    text = str(value).strip()
    for fmt in ("%d/%m/%Y", "%Y-%m-%d", "%d %b %Y", "%d %B %Y", "%d-%m-%Y", "%m/%d/%Y"):
        try:
            return datetime.strptime(text, fmt).date().isoformat()
        except ValueError:
            continue
    return None


def parse(
    payload: bytes,
    lookback_days: int | None = None,
    first_appointments_only: bool = True,
) -> list[Matter]:
    """Read the Data set sheet into Matters.

    The sheet holds the full history - over 70,000 appointments back to July
    2021 - so `lookback_days` is what keeps a weekly run to the new ones. Pass
    0 to read everything.

    `first_appointments_only` keeps just the rows flagged Series 1, a company's
    first entry into external administration. Series 2 rows repeat a company
    for each subsequent appointment, which would create duplicate matters and
    double-count the same creditors.
    """
    from openpyxl import load_workbook

    workbook = load_workbook(io.BytesIO(payload), read_only=True, data_only=True)
    try:
        sheet = _find_sheet(workbook)
        rows = list(sheet.iter_rows(values_only=True))
    finally:
        workbook.close()

    header_index, mapping = _find_header(rows)
    log.info(
        "Data set sheet: header on row %d, %d columns mapped: %s",
        header_index + 1, len(mapping), sorted(set(mapping.values())),
    )

    days = lookback_days
    if days is None:
        days = config.settings()["sources"]["asic_stats"]["lookback_days"]
    cutoff = (date.today() - timedelta(days=days)).isoformat() if days else None

    matters: list[Matter] = []
    skipped_old = skipped_repeat = skipped_solvent = 0
    newest = ""
    for row in rows[header_index + 1:]:
        values: dict[str, Any] = {}
        for column, field in mapping.items():
            if column < len(row):
                values[field] = row[column]

        company = str(values.get("company_name") or "").strip()
        if not company or _norm_header(company) in HEADER_ALIASES:
            continue

        # Series 1 is only present when the mapping found the column; if ASIC
        # drops it, fall back to keeping every row rather than keeping none.
        track_series_1 = first_appointments_only and "series_1" in mapping.values()
        if track_series_1 and not _is_true(values.get("series_1")):
            skipped_repeat += 1
            continue

        appointment_type = str(values.get("appointment_type") or "").strip() or None
        if appointment_type and any(
            s in appointment_type.lower() for s in SOLVENT_APPOINTMENTS
        ):
            skipped_solvent += 1
            continue

        appointment_date = _parse_date(values.get("appointment_date"))
        if appointment_date and appointment_date > newest:
            newest = appointment_date
        if cutoff and appointment_date and appointment_date < cutoff:
            skipped_old += 1
            continue

        matters.append(
            Matter(
                source="asic",
                company_name=company,
                acn=_normalise_acn(values.get("acn")),
                appointment_type=appointment_type,
                appointment_date=appointment_date,
                practitioner=str(values.get("practitioner") or "").strip() or None,
                practitioner_firm=str(values.get("practitioner_firm") or "").strip() or None,
                industry=str(values.get("industry") or "").strip() or None,
                industry_subdivision=str(values.get("industry_subdivision") or "").strip()
                or None,
                state=str(values.get("state") or "").strip() or None,
                postcode=str(values.get("postcode") or "").strip() or None,
                first_seen=date.today().isoformat(),
            )
        )

    log.info(
        "Data set sheet: %d appointments in the last %s days "
        "(skipped %d older, %d repeat appointments, %d solvent wind-ups)",
        len(matters), days or "all", skipped_old, skipped_repeat, skipped_solvent,
    )

    # The workbook trails reality by weeks. If the window cannot reach the
    # newest row in the file, the run finds nothing and looks like a quiet
    # week - so say exactly what happened instead.
    if cutoff and not matters and newest:
        lag = (date.today() - date.fromisoformat(newest)).days
        log.error(
            "No matters returned: the newest appointment in the file is %s "
            "(%d days old) but the lookback window is only %d days. Widen "
            "sources.asic_stats.lookback_days.",
            newest, lag, days,
        )
    if not matters and not (skipped_old or skipped_repeat or skipped_solvent):
        raise RuntimeError(
            "The data set sheet produced no rows at all. The schema has probably "
            "changed - run `python -m creditor_sourcing schema`."
        )
    return matters


def collect(client: Client | None = None, url: str | None = None) -> list[Matter]:
    return parse(download(url, client))


def describe(payload: bytes, sample_rows: int = 6) -> str:
    """Human-readable dump of the workbook's shape, for calibration.

    Printed by `python -m creditor_sourcing schema`. This is how the real
    column names get discovered from an environment that cannot reach ASIC.
    """
    from openpyxl import load_workbook

    workbook = load_workbook(io.BytesIO(payload), read_only=True, data_only=True)
    lines: list[str] = [f"Sheets: {workbook.sheetnames}", ""]
    try:
        sheet = _find_sheet(workbook)
        lines.append(f"Selected sheet: {sheet.title!r}")
        rows = []
        for index, row in enumerate(sheet.iter_rows(values_only=True)):
            rows.append(row)
            if index >= 30:
                break
        for index, row in enumerate(rows[: sample_rows + 12]):
            cells = [str(c)[:28] if c is not None else "" for c in row[:14]]
            lines.append(f"  row {index + 1:>3}: {cells}")
        try:
            header_index, mapping = _find_header(rows)
            lines += [
                "",
                f"Header row detected: {header_index + 1}",
                "Mapped columns:",
                *[
                    f"  col {column + 1:>2} -> {field}"
                    for column, field in sorted(mapping.items())
                ],
                "",
                "Unmapped headers (add to HEADER_ALIASES if useful):",
                *[
                    f"  col {column + 1:>2}: {value!r}"
                    for column, value in enumerate(rows[header_index])
                    if column not in mapping and value not in (None, "")
                ],
            ]
        except RuntimeError as exc:
            lines += ["", f"HEADER DETECTION FAILED: {exc}"]
    finally:
        workbook.close()
    return "\n".join(lines)


def save(payload: bytes, path: Path) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_bytes(payload)
    return path
