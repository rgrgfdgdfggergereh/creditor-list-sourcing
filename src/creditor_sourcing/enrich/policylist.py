"""Load the NCI PolicyList export so existing clients can be excluded.

The export arrives in three shapes and the loader has to take all of them:

* the NCI system export, an ``.xls`` whose header is on the SECOND row (row
  one is a report title);
* the same thing re-saved as ``.xlsx``;
* the ``.csv`` that the bulk-prospecting skill carries, header on row one.

Two columns name a company. ``Policy Name`` is the insured entity on the
schedule and ``Client Name`` is the account it sits under - and they differ
for group accounts: the policy ``MITRE 10 AUSTRALIA PTY LTD`` sits under the
client ``TOTAL TOOLS & HARDWARE GROUP``. Reading only one column let Mitre 10
and Home Timber & Hardware through as fresh prospects in the 13 September
list, so every name-bearing column is loaded.
"""

from __future__ import annotations

import csv
import logging
from collections.abc import Iterable, Iterator, Sequence
from pathlib import Path
from typing import Any

from ..models import normalise_name

log = logging.getLogger(__name__)

# Header text that marks a column holding a company name. Compared against
# the lower-cased header with underscores folded to spaces, first for
# equality, then as a fragment, so ``client_name``, ``Client Name`` and
# ``Policy Holder Name`` all qualify.
NAME_HEADERS = (
    "client name", "policy name", "company name", "insured", "insured name",
    "policy holder", "policyholder", "policy holder name", "name", "client",
)

# How far down to look for the header row. The NCI export puts it on row 2.
HEADER_SCAN_ROWS = 5


def _rows(path: Path) -> Iterator[Sequence[Any]]:
    suffix = path.suffix.lower()
    if suffix == ".csv":
        with path.open(encoding="utf-8-sig", newline="") as fh:
            yield from csv.reader(fh)
        return
    if suffix == ".xls":
        import xlrd

        book = xlrd.open_workbook(str(path))
        sheet = book.sheet_by_index(0)
        for i in range(sheet.nrows):
            yield [cell.value for cell in sheet.row(i)]
        return
    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        yield from workbook.worksheets[0].iter_rows(values_only=True)
    finally:
        workbook.close()


def _fold(header: Any) -> str:
    return str(header or "").strip().lower().replace("_", " ")


def name_columns(header: Sequence[Any]) -> list[int]:
    """Indexes of every column in ``header`` that holds a company name."""
    folded = [_fold(h) for h in header]
    exact = [i for i, h in enumerate(folded) if h in NAME_HEADERS]
    if exact:
        return exact
    return [
        i for i, h in enumerate(folded)
        if h and any(n in h for n in NAME_HEADERS)
        and not any(bad in h for bad in ("contact", "first", "surname", "email"))
    ]


def _find_header(rows: Iterator[Sequence[Any]]) -> tuple[Sequence[Any], list[int]] | None:
    for _ in range(HEADER_SCAN_ROWS):
        row = next(rows, None)
        if row is None:
            return None
        columns = name_columns(row)
        if columns:
            return row, columns
    return None


def load(path: Path) -> dict[str, str]:
    """Return {normalised name: original name} for every policyholder.

    Names from every name-bearing column are folded into one lookup, so a
    prospect matches whether it appears as the client or as the insured
    entity on a policy under a different client.
    """
    if not path.exists():
        log.warning("PolicyList not found at %s - client exclusion is OFF", path)
        return {}

    rows = _rows(path)
    found = _find_header(rows)
    if found is None:
        log.warning(
            "No client-name column in the first %d rows of %s - client exclusion is OFF",
            HEADER_SCAN_ROWS, path,
        )
        return {}
    header, columns = found

    clients = fold_names(
        (row[c] for row in rows for c in columns if c < len(row))
    )
    log.info(
        "PolicyList: %d client names loaded from %s (columns: %s)",
        len(clients), path.name,
        ", ".join(str(header[c]) for c in columns),
    )
    return clients


def fold_names(names: Iterable[Any]) -> dict[str, str]:
    """Normalise names into the lookup shape ``policylist_match`` consumes."""
    clients: dict[str, str] = {}
    for raw in names:
        name = str(raw or "").strip()
        key = normalise_name(name)
        if key:
            clients.setdefault(key, name)
    return clients
